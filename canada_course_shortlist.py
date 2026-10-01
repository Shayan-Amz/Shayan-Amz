#!/usr/bin/env python3
"""Build Shayan's security-free, course/project-based Canadian shortlist.

Rankings and academic rules are sourced. Admission ranges are MANUAL, SUBJECTIVE,
UNCALIBRATED planning judgements, NOT measured acceptance probabilities or CIs.
No admissions model is trained and no migration-country score is used here.
Existing migration deliverables and the original transcript are never modified.
"""
from __future__ import annotations

import csv
import argparse
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from math import inf
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent
DATE = "2026-10-01"
FX = Decimal("1.419")  # Same USD/CAD comparison assumption as the earlier workbook; not live FX.
PROOF_CAD = Decimal("23448")  # IRCC, effective September 1, 2026; 2027 must be rechecked.
BUDGET_USD = 45000
IGNORE_BUDGET = False
OUTPUT_STEM = "Canada-Course-Shortlist"
FONT = "B Nazanin"

SOURCES = {
    "rank_cs": ("THE Computer Science 2026 — جدول رسمی کانادا", "https://www.timeshighereducation.com/student/best-universities/best-universities-canada-computer-science-degrees"),
    "rank_uni": ("QS — رتبه‌های کلی جاری کانادا، نسخهٔ ۲۰۲۷", "https://www.topuniversities.com/where-to-study/north-america/canada/guide"),
    "proof": ("IRCC — تمکن، از سپتامبر ۲۰۲۶؛ بدون اتکا به کار دانشجویی", "https://www.canada.ca/en/immigration-refugees-citizenship/services/study-canada/study-permit/get-documents/financial-support.html"),
    "mun_fees": ("Memorial — جدول هزینهٔ رسمی 2026/27", "https://www.mun.ca/sgs/media/production/memorial/academic/school-of-graduate-studies/school-of-graduate-studies/media-library/MinimumExpense.pdf"),
    "mun_soft": ("Memorial — MASc Software Engineering؛ پذیرش محدود و رقابتی", "https://www.mun.ca/university-calendar/school-of-graduate-studies/school-of-graduate-studies/8/7/"),
    "mun_cs": ("Memorial — MSc CS، مسیر درسی و پایان‌نامه‌ای", "https://www.mun.ca/become/graduate/programs-and-courses/computer-science/"),
    "mun_ai": ("Memorial — MAI، پیش‌نیازها، درسی، بدون co-op داخلی", "https://www.mun.ca/become/graduate/programs-and-courses/artificial-intelligence/"),
    "mun_ds": ("Memorial — MDSc، پیش‌نیازهای کمی و دوره‌های آمادگی", "https://www.mun.ca/become/graduate/programs-and-courses/data-science/"),
    "mcm_soft": ("McMaster — MEng Computing and Software", "https://www.eng.mcmaster.ca/cas/degree-options/computing-and-software-meng/"),
    "mcm_fees": ("McMaster — لینک جدول رسمی 2026/27؛ مقدار برنامه در خوانش فعلی استخراج نشد", "https://registrar.mcmaster.ca/fees/graduate/"),
    "queen": ("Queen's — MEng ECE؛ حد ایران ۱۵/۲۰، IELTS 7، internship فعلاً عرضه نمی‌شود", "https://smithengineering.queensu.ca/ece/graduate/meng"),
    "calg_soft": ("Calgary — MEng Software، ۱۰/۱۳ درس، سابقهٔ نرم‌افزاری و مهلت ویژهٔ ایران", "https://grad.ucalgary.ca/future-students/explore-programs/electrical-and-computer-engineering-meng-software-course"),
    "regina": ("Regina — MEng SSE، درسی/پروژه‌ای؛ با گواهی AI تعلیق‌شده متفاوت است", "https://www.uregina.ca/graduate-studies-research/graduate-calendar/all-programs/engg-sse.html"),
    "regina_wes": ("Regina Engineering — ارزیابی WES برای مدارک بین‌المللی", "https://www.uregina.ca/engineering/graduate-students/Index.html"),
    "ont_soft": ("Ontario Tech — MEng Software؛ B کل و دو سال آخر", "https://gradstudies.ontariotechu.ca/future_students/programs/masters_programs/software_engineering/index.php"),
    "uvic_project": ("UVic — MEng ECE، پروژه و استاد الزامی؛ درسیِ صرف نیست", "https://www.uvic.ca/graduate/programs/graduate-programs/credential-pages/electrical-computer-engineering-cred/electrical-and-computer-engineering-meng.php"),
    "windsor": ("Windsor — MAC؛ سه سال آخر، نمرات CS و زبان", "https://www.future.uwindsor.ca/program/master-of-applied-computing/"),
    "tmu_cs": ("TMU — MSc CS، شرط مسیر Course Option", "https://www.torontomu.ca/cs/graduate/admissions-applying/"),
    "tmu_cs_degree": ("TMU — گزینهٔ درسی MSc CS، ۸ درس و زمان استاندارد", "https://www.torontomu.ca/cs/graduate/degree-requirements/"),
    "unb": ("UNB — MCSC، درسی، ۱۰ درس و co-op اختیاری", "https://www.unb.ca/fredericton/cs/grad/masters/mcsc.html"),
    "unb_entry": ("UNB — حد B برای مسیر درسی، زمینهٔ CS و رقابت پذیرش", "https://www.unb.ca/fredericton/cs/grad/admission.html"),
    "manitoba": ("Manitoba — MEng ECE، صنعتی/پروژه‌ای و پذیرش استاد", "https://umanitoba.ca/graduate-studies/admissions/programs-of-study/electrical-and-computer-engineering-msc-meng"),
    "ottawa_ai": ("Ottawa — MEng ECE Applied AI، انگلیسی، B+ و پروژه", "https://catalogue.uottawa.ca/en/graduate/master-engineering-electrical-computer-engineering-concentration-applied-artificial-intelligence/"),
    "york_ai": ("York — MSc CS AI، پروژه‌ای و بدون بستهٔ فاند", "https://lassonde.yorku.ca/eecs/academics/specialization-in-artificial-intelligence/"),
    "uvic_ds": ("UVic — MEng Applied Data Science، یک‌ساله، درسی و self-funded", "https://www.uvic.ca/graduate/programs/graduate-programs/credential-pages/electrical-computer-engineering-cred/applied-data-science-meng.php"),
    "tmu_ds": ("TMU — Data Science MRP؛ تحویل مجازی/ترکیبی و پیش‌نیازها", "https://www.torontomu.ca/graduate/programs/data-science-analytics/"),
    "pgwp": ("IRCC — PGWP؛ حداقل ۵۰٪ حضوری در کانادا برای دوره‌های جدید", "https://www.canada.ca/en/immigration-refugees-citizenship/services/study-canada/work/after-graduation/eligibility.html"),
    "alberta_mm": ("Alberta Multimedia — کمیتهٔ مستقل، GPA دو سال آخر و برنامه‌نویسی", "https://mmgrad.org/program.php"),
    "alberta_fee_method": ("Alberta Multimedia — هزینهٔ تمام ۳۶ واحد، برنامه self-funded", "https://mmgrad.org/courses.php"),
    "alberta_nonstandard": ("Alberta — جدول رسمی Non-standard 2026/27، از صفحهٔ دانشگاه", "https://www.ualberta.ca/en/graduate-studies/fees-funding/tuition-fees/non-standard.html"),
    "alberta_rate_doc": ("Alberta — نرخ Multimedia بین‌المللی 2026/27، هر سه واحد", "https://docs.google.com/document/d/1N-R4kxLGsynAxLxRgk6NYbSEzR_99L-vyOFasYb9z0Q/export?format=html"),
    "tmu_media": ("TMU — Master of Digital Media؛ ۱۲ ماه، B دو سال آخر و portfolio", "https://www.torontomu.ca/graduate/programs/digital-media/"),
    "waterloo_meng": ("Waterloo — MEng ECE، نرم‌افزار، درسی و بدون استاد برای ورود", "https://uwaterloo.ca/future-graduate-students/programs/by-faculty/engineering/electrical-and-computer-engineering-master-engineering-meng"),
    "waterloo_fees": ("Waterloo — شهریهٔ بین‌المللی رسمی Fall 2026", "https://uwaterloo.ca/finance/masters-and-phd-program-tuition-fall-2026-international-0"),
    "western_meng": ("Western — MEng ECE؛ ۳ ترم، حد ۷۰٪ دو سال آخر و توجه به کار", "https://grad.uwo.ca/admissions/program.cfm?p=40"),
    "western_soft": ("Western — گرایش Software در MEng", "https://www.eng.uwo.ca/electrical/graduate/current_students/meng_programs/index.html"),
    "western_fees": ("Western — جدول بین‌المللی Graduate، Fall 2026، MEng ECE", "https://www.registrar.uwo.ca/student_finances/fees_refunds/Fall-2026-Graduate-Fee-Schedules---FT-International.pdf"),
    "waterloo_ds": ("Waterloo — Data Science؛ حد واقعی پذیرش بسیار بالاتر از حداقل", "https://uwaterloo.ca/data-science/graduate-programs/admission-requirements"),
    "waterloo_ds_lang": ("Waterloo — MDSAI IELTS 7.5 و W/S حداقل 7", "https://uwaterloo.ca/data-science/graduate-programs/frequently-asked-questions-faqs"),
    "western_mda": ("Western — MDA؛ حد ۷۵٪ در تک‌تک پیش‌نیازهای کمی", "https://www.uwo.ca/mda/admissions/index.html"),
    "toronto_mscac": ("Toronto — MScAC؛ قواعد رسمی و چرخهٔ Fall 2027: 1 Oct تا 1 Dec 2026", "https://mscac.utoronto.ca/apply/"),
    "toronto_mscac_sgs": ("Toronto SGS — B+ سال آخر، گرایش‌ها و applied-research internship", "https://www.sgs.utoronto.ca/programs/applied-computing/"),
    "ubc": ("UBC — ایران، معدل کل ۱۶/۲۰", "https://www.grad.ubc.ca/country/iran"),
    "sfu": ("SFU — ایران، معدل کل ۱۶/۲۰", "https://www.sfu.ca/gradstudies/apply/applying/requirements/iran.html"),
    "alberta_closed": ("Alberta — MSc درسی عمومی CS متوقف است؛ Multimedia جداست", "https://www.ualberta.ca/en/computing-science/graduate-studies/programs-and-admissions/applications-and-admissions/index.html"),
    "capstone": ("Portfolio — README پروژهٔ کارشناسی املاک؛ ادعای استقرار، نه ممیزی live", "https://github.com/Shayan-Amz/Real-Estate-File-Manager-Project"),
    "pipeline": ("Portfolio — پایپ‌لاین داده و پیشنهادگر تطبیق ژانر، نه آموزش مدل", "https://github.com/Shayan-Amz/Movie-ML-Pipeline"),
    "unity": ("Portfolio — تمرین Unity/C# UI، نه شاهد پژوهش ML", "https://github.com/Shayan-Amz/Unity-School-Directory-UI"),
}

# Same ranking publisher/year for every subject comparison. Bands are not broken
# into fictitious exact positions. The rank is the university's CS field, NOT a
# ranking of a named MSc/MEng, Software Engineering, Multimedia or Digital Media.
UNI = {
    "Waterloo": ("41", "113"), "Alberta": ("77", "96"),
    "Western": ("201–250", "142"), "Ottawa": ("201–250", "228"),
    "York": ("201–250", "322"), "Victoria": ("201–250", "370"),
    "McMaster": ("251–300", "174"), "Memorial": ("251–300", "721–730"),
    "Queen's": ("301–400", "179"), "Calgary": ("301–400", "249"),
    "Windsor": ("301–400", "537"), "TMU": ("301–400", "669"),
    "UNB": ("301–400", "677"), "Regina": ("401–500", "1001–1200"),
    "Ontario Tech": ("601–800", None), "Manitoba": (None, "560"),
    "Toronto": ("22", "32"),
}

GROUPS = ["نرم‌افزار", "کامپیوتر کاربردی", "هوش مصنوعی و داده", "رسانه و تعامل"]
ITEMS = []
def add(group, uni, program, route, chance, entry, why, sources,
        fee_cad=None, budget_note="شهریهٔ کامل بین‌المللی ۲۰۲۷ هنوز تأیید نشده؛ داخل بودجه فرض نشود.", priority="مشروط به بررسی", duration="—"):
    assert group in GROUPS
    ITEMS.append(dict(group=group, uni=uni, program=program, route=route,
        chance=chance, entry=entry, why=why, sources=sources, fee_cad=fee_cad,
        budget_note=budget_note, priority=priority, duration=duration))

add("نرم‌افزار", "McMaster", "MEng Computing and Software", "درسی + پروژه", (20,40),
    "کارشناسی مرتبط؛ حد اعلام‌شده B−؛ توصیه‌نامه/پروژه و معادل‌سازی لازم.",
    "نرم‌افزار ۱۹٫۲۵، پروژه ۱۹٫۵ و برنامهٔ املاک به نفع تطابق‌اند؛ الگوریتم ۱۱٫۵ ریسک باقی می‌ماند.",
    ["mcm_soft","mcm_fees"], priority="بررسی جدی علمی؛ هزینه نامعلوم", duration="حدود ۲ سال")
add("نرم‌افزار", "Memorial", "MASc Software Engineering", "درسی + capstone", (30,55),
    "حداقل second-class کارشناسی چهار‌سالهٔ CS/CE؛ پذیرش محدود و رقابتی؛ زبان بالاترِ SGS.",
    "بیشترین تطابق مستقیم با نرم‌افزار، OS، پروژهٔ ۱۹٫۵ و تجربهٔ اجرایی؛ درس Applied Algorithms دارد و باید ضعف الگوریتم را جدی گرفت.",
    ["mun_soft","mun_fees"], fee_cad=Decimal("29948"),
    budget_note="نرخ 2026/27؛ گام۱ پایه زیر سقف است، ولی ancillary، سفر/ویزا و نرخ ۲۰۲۷ اضافه‌اند.",
    priority="اولویت بررسی", duration="۱۶ ماه")
add("نرم‌افزار", "Queen's", "MEng ECE — Computer and Software Engineering", "درسی؛ پروژه اختیاری", (20,40),
    "حد ایران ۱۵/۲۰؛ IELTS حداقل ۷؛ استاد برای پذیرش لازم نیست.",
    "مدرک CE و درس‌های نرم‌افزاری متناسب‌اند؛ internship در 2026/27 عرضه نمی‌شود و بازگشتش برای ۲۰۲۷ تضمین نیست.",
    ["queen"], duration="۸ ماه؛ پروژه/تأخیر می‌تواند طول را بیشتر کند")
add("نرم‌افزار", "Calgary", "MEng ECE — Software Engineering", "درسی؛ پروژهٔ تیمی", (25,45),
    "GPA معادل ۳٫۰ در دو سال آخر؛ کارشناسی مهندسی؛ برای شروع سپتامبر، سابقهٔ آموزشی Software باید تأیید شود.",
    "دروس نرم‌افزار و برنامهٔ املاک مفیدند؛ تأیید ۱۰درس/معافیت از ۳ درس foundation لازم است، نه صرف فرض از نام CE. برای ایران مهلت ویژه در صفحه آمده؛ چرخهٔ ۲۰۲۷ تأیید شود.",
    ["calg_soft"], duration="۸/۱۲ ماه فشرده یا تا ۲ سال")
add("نرم‌افزار", "Regina", "MEng Software Systems Engineering", "درسی + پروژه", (30,55),
    "معادل‌سازی حداقل مهندسی؛ ارزیابی WES برای مدرک بین‌المللی؛ پروژه دارای supervisor.",
    "برای software systems، backend و پروژهٔ کاربردی متناسب است. رنک کلی پایین‌تر است؛ شانس بالا/قطعی از آن نتیجه نمی‌شود. گواهی AI تعلیق‌شده را با این مدرک اشتباه نگیریم.",
    ["regina","regina_wes"], duration="حدود ۲ سال")
add("نرم‌افزار", "Ontario Tech", "MEng Software Engineering", "درسی؛ امکان انتقال به پروژه", (25,45),
    "حداقل B در کل و دو سال آخر؛ مدرک مهندسی مرتبط؛ معادل دقیق ایران باید تأیید شود.",
    "نرم‌افزار و capstone مناسب‌اند. Advanced Engineering Mathematics جزو core است؛ این مسیر کاملاً بدون ریاضی نیست. نرخ شهریهٔ MITS را برای MEng به کار نبرده‌ام.",
    ["ont_soft"], duration="۱۶ ماه")

add("کامپیوتر کاربردی", "Victoria", "MEng Electrical and Computer Engineering — software/systems project", "پروژه‌ای؛ درسیِ صرف نیست", None,
    "معادل حداقل FGS؛ استاد باید پذیرش را توصیه کند؛ نام استاد روی درخواست لازم است.",
    "OS، زبان‌ها، معماری و پروژه مناسب‌اند. بدون دیدن امکان گرفتن استاد/پروژه، درصد نمی‌سازم؛ بعد از موافقت استاد نیز تصمیم نهایی دانشگاه مستقل است.",
    ["uvic_project"], duration="۲ سال", priority="استادمحور؛ پشتیبان پروژه‌ای")
add("کامپیوتر کاربردی", "Memorial", "MSc Computer Science — all-course route", "درسی؛ نه thesis", (25,45),
    "کارشناسی CS/CE؛ حد ۷۵٪ کل یا معادل؛ GRE برای بین‌المللی توصیه شده، نه جایگزین معدل.",
    "نرم‌افزار/OS خوب است و مسیر اقتصادی است؛ اما CS عمومی الگوریتم/نظریه هم می‌خواهد و از MASc نرم‌افزار برایت کم‌تطابق‌تر است.",
    ["mun_cs","mun_fees"], fee_cad=Decimal("15600"),
    budget_note="شهریهٔ ۶ پرداختِ مسیر درسی 2026/27؛ حتی پایان زودتر به معنی حذف پرداخت‌های لازم نیست. هزینه‌های اجباری/۲۰۲۷ جدا.",
    priority="اولویت بررسی اقتصادی", duration="۲ سال")
add("کامپیوتر کاربردی", "Windsor", "Master of Applied Computing (MAC)", "درسی/کاربردی", (20,40),
    "حد major average یا سه سال آخر در صفحهٔ MAC؛ IELTS ۶٫۵ با قیود band؛ قواعد CS پایین و معادل‌سازی لازم.",
    "پروژه و برنامه‌نویسی مثبت‌اند؛ معدل سه سال آخر ۱۶٫۱۹ است. چند نمرهٔ پایین CS مانع برچسب Safe می‌شود؛ raw ۱۴/۲۰ را خودکار معادل ۷۰٪ نمی‌دانم.",
    ["windsor"], duration="مدت و مسیر internship نهایی از برنامه تأیید شود")
add("کامپیوتر کاربردی", "TMU", "MSc Computer Science — Course Only", "درسی", (20,40),
    "حد B / ۳٫۰۰ از ۴٫۳۳ یا معادل، دو سال آخر؛ دو توصیه‌نامه از استادان؛ مسیر Course انتخاب شود.",
    "دروس و پروژهٔ software به نفع توست؛ الگوریتم پایین را پنهان نکنیم. سابقهٔ حرفه‌ای IT را سه‌سال توسعهٔ حرفه‌ای نرم‌افزار فرض نمی‌کنم.",
    ["tmu_cs","tmu_cs_degree"], duration="استاندارد ۲ سال")
add("کامپیوتر کاربردی", "UNB", "Master of Computer Science by Coursework (MCSC)", "درسی؛ co-op اختیاری", (25,50),
    "حد B برای درسی در صفحهٔ پذیرش؛ زمینهٔ software/data structures/OS/algorithms/architecture لازم؛ زبان احراز شود.",
    "برای IT professionals و مهارت‌های کاربردی طراحی شده؛ دروس قوی و پروژهٔ deployed متناسب‌اند، اما الگوریتم ریسک است. campus Fredericton، تمام‌وقت؛ Saint John را برای تمام‌وقت جایگزین نکرده‌ام.",
    ["unb","unb_entry"], duration="طول و ۱۰درس/دو co-op از برنامه تأیید شود")
add("کامپیوتر کاربردی", "Manitoba", "MEng Electrical and Computer Engineering — industrial project", "درسی + پروژه؛ نیازمند استاد", None,
    "معادل GPA حداقل ۳٫۰؛ پذیرش اولیهٔ academic advisor لازم؛ مدرک مرتبط.",
    "نرم‌افزار/سیستم‌ها می‌تواند موضوع پروژه باشد؛ اسم و موافقت استاد نداریم، پس درصد فردی نمی‌دهم. رتبهٔ CS در جدول مرجع یافت نشد؛ دانشگاه حذف نشده است.",
    ["manitoba"], duration="۲ سال", priority="استادمحور؛ پشتیبان پروژه‌ای")

add("هوش مصنوعی و داده", "Ottawa", "MEng ECE — Applied Artificial Intelligence", "درسی + research project؛ بدون thesis", (15,30),
    "مدرک ECE/معادل و حد B+ یا ۷۵٪ معادل؛ آموزش انگلیسی؛ پیش‌نیازها/زبان تأیید شود.",
    "پروژهٔ املاک با استفادهٔ کاربردی از AI مفید است؛ آمار ۱۲، حسابان ۱۰ و ML ۱۴٫۶۵ در این curriculum ریسک‌اند. تخفیفِ نیازمند فرانسوی در بودجه حساب نشده است.",
    ["ottawa_ai"], duration="تا ۲ سال")
add("هوش مصنوعی و داده", "York", "MSc Computer Science — AI Specialization", "پروژه‌ای؛ non-thesis", (20,40),
    "B+ دو سال آخر؛ درس senior-level theoretical CS؛ IELTS ۷؛ GRE توصیهٔ قویِ صفحه، نه تضمین جبران معدل.",
    "AI کاربردیِ پروژه‌ای با capstone/ادغام AI تو مرتبط است؛ ریاضی و آمار ریسک‌اند. این MSc فنی با MMAI مدیریتیِ Schulich فرق دارد و بستهٔ فاند ندارد؛ internship را دانشجو پیدا می‌کند.",
    ["york_ai"], duration="حدود ۲۰ ماه")
add("هوش مصنوعی و داده", "Victoria", "MEng Applied Data Science (MADS)", "درسی؛ co-op اختیاری پس از coursework", (15,30),
    "مدرک CE/CS/مرتبط، حداقل عمومی FGS، دو assessment report، رزومه و statement.",
    "پایپ‌لاین داده و برنامه‌نویسی مرتبط‌اند؛ ضعف آمار و حسابان همچنان مهم است. رقم رسمی کامل برای ۲۰۲۷ استخراج نشده؛ بر اساس aggregator زیر سقف اعلام نشده است.",
    ["uvic_ds"], duration="۱ سال")
add("هوش مصنوعی و داده", "Memorial", "Master of Artificial Intelligence (MAI)", "درسی + capstone", (20,40),
    "Second-Class Upper / ۷۵٪ معادل؛ جبر خطی، حسابان برداری، آمار و برنامه‌نویسی؛ زبان بالاترِ SGS.",
    "AI ۱۷٫۲۵، جبر ۱۸ و پروژهٔ AI-integrated مثبت‌اند؛ ML و حسابان/آمار قوی نیستند. ادغام Whisper/LLM، پژوهش یا آموزش مدل را ثابت نمی‌کند. داخل برنامه co-op/بورس اعلام نشده.",
    ["mun_ai","mun_fees"], fee_cad=Decimal("29948"),
    budget_note="پایهٔ 2026/27 شامل regular + special fee؛ ancillary و هزینه‌های ۲۰۲۷ جدا.", duration="۱۶ ماه")
add("هوش مصنوعی و داده", "Memorial", "Master of Data Science (MDSc)", "درسی + capstone", (15,30),
    "second-class؛ حسابان چندمتغیره، استنباط آماری و programming یا دانش معادل؛ دوره‌های آمادگی قبل شروع.",
    "ETL و Python مفیدند؛ پیشنهادگرِ پروژهٔ سینما rule-based است، نه مدل ML آموزش‌دیده. پوشش inference از اسم درس آمار قطعی نیست و نمره ۱۲ ضعف است.",
    ["mun_ds","mun_fees"], fee_cad=Decimal("26936"),
    budget_note="پایهٔ 2026/27، Plan C و special fee؛ ancillary/۲۰۲۷ جدا. قیمت ارزان‌تر به معنی تناسب علمی بالاتر نیست.", duration="۱ سال")
add("هوش مصنوعی و داده", "TMU", "MSc Data Science and Analytics — MRP option", "درسی + major research project", (15,30),
    "B دو سال آخر؛ دانش آمار، داده‌ساختار، پایگاه داده و R/Python؛ دو reference.",
    "Python/ETL و نرم‌افزار مثبت، آمار/حسابان ضعیف. تحویل virtual/hybrid دارد؛ پیش از هزینه، امکان حداقل ۵۰٪ حضوری در کانادا و PGWP باید کتبی تأیید شود.",
    ["tmu_ds","pgwp"], duration="۱ سال تمام‌وقت", priority="مشروط به حضوری/PGWP و هزینه")

add("رسانه و تعامل", "Alberta", "MSc Computing Science — Multimedia", "درسی + R&D internship", (20,40),
    "حد ۳٫۰/۴٫۰ معادل در دو سال آخر؛ کمیتهٔ مستقل؛ programming جدی؛ تجربهٔ صنعتی ممکن است لحاظ شود.",
    "C#/C++، Unity، سیستم تعاملی و پروژهٔ کاربردی مثبت‌اند؛ گرافیک ۱۳٫۷۵/حسابان پایین ریسک‌اند. internship بخشی از دوره است ولی حقوق/محل صنعتی را تضمین نکرده‌ام.",
    ["alberta_mm","alberta_fee_method","alberta_nonstandard","alberta_rate_doc"],
    fee_cad=Decimal("2659.8")*12,
    budget_note="نرخ رسمی 2026/27 برای تمام ۳۶ واحد: ۱۲ سهم سه‌واحدی؛ نرخ قدیمی ۲۰۲۲ به کار نرفته. ancillary/۲۰۲۷ جدا.",
    priority="بررسی جدی با portfolio مرتبط", duration="طراحی‌شده برای حدود ۲ سال؛ ۴ term ثبتِ ۹واحدی")
add("رسانه و تعامل", "TMU", "Master of Digital Media (MDM)", "حرفه‌ای/درسی و پروژه‌محور", (25,50),
    "حد B دو سال آخر؛ portfolio، رزومه، statement و دو/سه توصیه‌نامه.",
    "Unity، UI/UX و محصول املاک متناسب‌اند؛ این برنامه CS/Software Engineering خالص نیست و جنبهٔ طراحی/محصول/کارآفرینی دارد. رتبه CS صرفاً proxy مؤسسه است، نه رتبه Digital Media.",
    ["tmu_media"], duration="۱۲ ماه")

EXCLUDED = [
    dict(uni="Waterloo", program="MEng ECE — Software", chance=(10,25), fee_cad=Decimal("17400")*4,
         reason="تناسب نرم‌افزاری دارد، اما ۴ term استاندارد شهریهٔ فعلی، گام۱ را حدود $65.6k می‌کند؛ حتی ۳ term حدود $53.3k قبل ancillary است. کوتاه‌کردن به ۲ term بدون تأیید رسمی بار درس/هزینه مبنا نیست.", sources=["waterloo_meng","waterloo_fees"]),
    dict(uni="Western", program="MEng ECE — Software", chance=(25,45), fee_cad=Decimal("15717.33")*3,
         reason="از نظر تحصیلی/کار مرتبط است؛ ۳ term نرخ Fall 2026، گام۱ پایه حدود $49.8k و قبل ancillary است؛ بدون کاهش رسمی هزینه/بورس از سقف عبور می‌کند.", sources=["western_meng","western_soft","western_fees"]),
    dict(uni="Waterloo", program="MDSAI", chance=None, fee_cad=Decimal("6272")*9,
         reason="گام۱ حدود $56.3k قبل ancillary/co-op؛ IELTS 7.5 با W/S حداقل 7 می‌خواهد و در سناریوی IELTS 7 این شرط احراز نیست؛ ضعف کمی و رقابت شدید هم مستقل‌اند.", sources=["waterloo_ds","waterloo_ds_lang","waterloo_fees"]),
    dict(uni="Toronto", program="MScAC — AI / Data Science", chance=None, fee_cad=Decimal("90050"),
         reason="شهریه و fees برآورد رسمی ۲۰۲۶ به‌تنهایی بالاتر از کل بودجه است؛ گام۱ نزدیک $80k. نام دانشگاه خوب، توان تأمین این فاصله را ایجاد نمی‌کند.", sources=["toronto_mscac"]),
    dict(uni="Western", program="Master of Data Analytics — AI", chance=None, fee_cad=None,
         reason="حد ۷۵٪ معادل در تک‌تک پیش‌نیازهای calculus/probability/statistics/linear algebra دارد. حسابان ۱۰ و آمار ۱۲ نیازمند تأییدِ احراز پیش‌نیازند؛ قبل آن درصد قابل دفاع نداریم.", sources=["western_mda"]),
    dict(uni="Alberta", program="General course-based MSc Computing Science", chance=None, fee_cad=None,
         reason="مسیر درسی عمومی از ۲۰۲۱ متوقف شده است؛ Multimedia مسیر جداگانه است.", sources=["alberta_closed"]),
]

ASSUMPTIONS = [
    ("دامنه", "در ادامهٔ بحث اخیر: فقط کانادا بدون کبک، ورودی سپتامبر ۲۰۲۷؛ تمام پیشنهادهای حوزهٔ امنیت حذف‌اند. نمرهٔ درس امنیت برای ساخت معدل رسمی پاک نشده است."),
    ("نوع مدرک", "درسی/پروژه‌ای اولویت دارد؛ thesis خالص در فهرست اصلی نیست. UVic/Manitoba MEng نیز استاد و پروژهٔ دفاع‌کردنی دارند، نه صرفاً کلاس."),
    ("معدل", "کل ۱۵٫۷۷؛ دو سال آخر ۱۶٫۹۲ روی ۷۷ واحد؛ سال آخر ۱۷٫۰۳ روی ۳۸ واحد؛ سه سال آخر ۱۶٫۱۹ روی ۱۱۱ واحد. از CSVِ استخراج‌شده دوباره محاسبه شده‌اند."),
    ("IELTS واقعی", "هنوز نمرهٔ واقعی ارائه نشده است؛ English course grade، IELTS یا معافیت زبان نیست."),
    ("سناریوی زبان", "برای برآوردهای شخصی: IELTS Academic overall 7.0 و هر band حداقل 6.5 فرض شده، نه احراز شده. اگر برنامه حد بالاتر بخواهد این سناریو برای آن کافی نیست."),
    ("سابقه کار", "گفتهٔ تو: حدود ۳ سال computer specialist و ۷ ماه Unity/C#. سه سال IT را سه سال توسعهٔ حرفه‌ای نرم‌افزار حساب نکرده‌ام. README شغل فعلی را از Mar 2025 ذکر می‌کند؛ سابقهٔ قبلی برای ادعای ۳ سال باید مستند شود و هم‌پوشانی مشاغل جمع ساده نشود."),
    ("portfolio", "README و metadata پروژه‌های public GitHub خوانده شد، نه ممیزی کامل code یا تأیید live deployment. capstone املاک، ETL سینما و Unity UI به نفع software/applications‌اند؛ exact-genre matching و API integration شاهد پژوهش/آموزش مدل ML نیستند."),
    ("SOP و توصیه‌نامه", "SOP اختصاصی و صادقانه؛ دو reference علمی قوی در برنامه‌های لازم و مدارک سابقه قابل اثبات فرض شده‌اند. متن این مدارک هنوز بررسی نشده است."),
    ("قانون درصدها", "درصدهای ستاره‌دار، فقط قضاوت شخصیِ خام و کالیبره‌نشده‌اند؛ نه آمار رسمی، نه مدل آموزش‌دیده، نه confidence interval، نه خروجی محاسبهٔ معتبر احتمال. دادهٔ پذیرفته/ردشدهٔ هم‌پروفایل نداریم. شانس واقعی ممکن است خارج این بازه‌ها باشد."),
    ("شرط اعتبارِ اعداد", "بازه‌های عددی فقط پس از احراز واقعی حداقل‌های دانشگاه (معادل GPA، مدرک، پیش‌نیاز، زبان) معنا دارند؛ وجود بازه احراز آن حداقل‌ها را ثابت نمی‌کند. اگر حداقل سخت احراز نشود، ابتدا همان مانع رفع شود."),
    ("موارد بدون درصد", "برای دو مسیرِ وابسته به استاد، موافقت و امکان پروژه را نمی‌دانیم؛ به‌جای عددسازی، درصد خالی گذاشته شده است. برای موارد دارای مانع سخت/زبان/بودجه هم درصد عمومی ساخته نشده."),
    ("نوع احتمال", "اگر از بازهٔ خام استفاده می‌کنی فقط برای مقایسه و تنوع سبد اپلای باشد؛ نه پیش‌بینی مالی. این بازه‌ها احتمال offer علمی‌اند، نه پذیرش با فاند، نه study permit، نه شغل، نه PR."),
    ("همبستگی اپلای‌ها", "قبولی دانشگاه‌ها مستقل نیست؛ از این درصدها برای محاسبهٔ احتمال حداقل یک قبولی با فرمول ضرب استفاده نکن."),
    ("رتبهٔ موضوعی", "THE Computer Science 2026، از جدول خود ناشر؛ رتبهٔ حوزهٔ CS در دانشگاه است، نه رتبهٔ خود MSc/MEng. Software/Applied Computing/AI/Multimedia رتبهٔ مستقل برنامه از این جدول ندارند؛ برای Digital Media شاخص CS فقط proxy مرتبطِ مؤسسه است."),
    ("رتبهٔ کلی", "QS World University Rankings 2027، از فهرست جاری خود QS. THE موضوعی و QS کلی دو شاخص متفاوت‌اند؛ عنوان و سال عمداً جدا نوشته شده‌اند."),
    ("مرتب‌سازی", "ابتدا حوزه؛ در هر حوزه رتبهٔ CS، سپس رتبهٔ کلی QS. رتبه‌های مشترک داخل band شکسته نشده‌اند. موضوعیِ نامعلوم آخر گروه است؛ برای نبود عدد، دانشگاه حذف نشده."),
    ("بودجه", "حدود $40,000؛ تا نزدیک $45,000 برای دانشگاه مناسب. بودجهٔ نامعلوم تأیید نشده است. ملاک مقایسهٔ این فایل فقط گام۱ پایه: کل شهریه + یک سال تمکن رسمی؛ ancillary، سفر، ویزا و زندگی واقعیِ باقی‌مانده هم لازم‌اند."),
    ("نرخ و سال هزینه", "فقط USD، با همان فرض مقایسهٔ قبلی: هر USD معادل 1.419 CAD؛ نرخ زنده یا پیش‌بینی ۲۰۲۷ نیست. شهریه‌های عددی از 2026/27‌اند؛ تمام rates و تمکن برای درخواست ۲۰۲۷ باید دوباره کنترل شوند."),
    ("تمکن و کار", "IRCC از ۱ سپتامبر ۲۰۲۶ تمکن پایهٔ یک نفر را اعلام کرده؛ معادل فرضی این فایل حدود $16,524 است. دولت هزینهٔ سال اول و برنامهٔ تأمین کل دوره را بدون اتکا به کار در کانادا می‌خواهد. درآمد co-op و کار دانشجویی تضمین تلقی نشده."),
    ("امنیت", "نبود علاقهٔ تو به امنیت پذیرفته شده؛ پایین‌بودن نمره را به فقدان توانایی تعبیر نمی‌کنم و هیچ برنامهٔ امنیتی پیشنهاد نمی‌دهم. معدلِ کارنامهٔ واقعی همچنان همهٔ نمره‌های مؤثر را دارد."),
    ("موارد تأییدنشده", "IELTS واقعی، تبدیل رسمی نمرات، محتوا/پیش‌نیاز بعضی درس‌ها، quality و ownership پروژه، reference و SOP، ظرفیت/قطع واقعی ۲۰۲۷، شهریهٔ بعضی برنامه‌ها و شرایط حضوری/PGWP."),
    ("سقف توانایی", "قضاوت تناسب کارنامه با برنامه است؛ نمرهٔ کم، سقف توانایی یادگیری آیندهٔ تو را تعیین نمی‌کند. IELTS 8 یا certificate غیرنمره‌دار، نمرهٔ تاریخی یا حداقل سخت را خودکار جبران نمی‌کند."),
]


def configure_scenario(ignore_budget=False):
    """Add the previously cost-excluded routes without changing academic gates.

    A fresh process uses the original 20-row budget-aware scenario by default.
    --ignore-budget produces separate outputs, preserving the original files.
    """
    global IGNORE_BUDGET, OUTPUT_STEM
    if not ignore_budget:
        return
    if IGNORE_BUDGET:
        return
    IGNORE_BUDGET = True
    OUTPUT_STEM = "Canada-Course-Shortlist-No-Budget"
    additions = [
        dict(group="نرم‌افزار", uni="Waterloo", program="MEng ECE — Software",
             route="درسی؛ پروژه اختیاری", chance=(10,25),
             entry="مدرک مرتبط؛ حد ۷۵٪ کل یا معادل برای بین‌المللی؛ زبان احراز شود؛ استاد برای ورود لازم نیست.",
             why="با حل فرضی بودجه به سبد برگشت. نرم‌افزار/OS/پروژه به نفع توست؛ معدل کل و الگوریتم پایین ریسک‌اند. حذف سقف مالی باعث بالا بردن بازهٔ قضاوتی قبلی نشده است.",
             sources=["waterloo_meng","waterloo_fees"], fee_cad=Decimal("17400")*4,
             budget_note="چهار term استاندارد با نرخ 2026/27؛ مبلغ صرفاً اطلاع است و عامل حذف نیست. ancillary و نرخ ۲۰۲۷ جدا.",
             priority="بررسی بلندپروازانهٔ نرم‌افزاری", duration="۱۶ ماه عادی", scenario_added=True),
        dict(group="نرم‌افزار", uni="Western", program="MEng ECE — Software",
             route="درسی یا پروژه‌ای", chance=(25,45),
             entry="حد ۷۰٪ معادل در دو سال آخر؛ تجربهٔ کار توجه ویژه دارد؛ گرایش Software و زبان تأیید شوند.",
             why="مانع قبلیِ حذف مالی برداشته شد؛ آخرین سال‌ها و درس‌های نرم‌افزار متناسب‌اند. این مسیر را با MDA آماریِ همین دانشگاه یکی ندانیم. بازهٔ قبلی فقط نظر خام است و با بودجه زیاد نشده.",
             sources=["western_meng","western_soft","western_fees"], fee_cad=Decimal("15717.33")*3,
             budget_note="سه term نرخ Fall 2026؛ صرفاً اطلاع هزینه، نه معیار حذف. پروژه/بورس تضمین نشده‌اند.",
             priority="بررسی جدیِ نرم‌افزاری", duration="۳ term / حدود ۱ سال", scenario_added=True),
        dict(group="هوش مصنوعی و داده", uni="Waterloo", program="MDSAI — Data Science and Artificial Intelligence",
             route="درسی + co-op؛ بدون thesis", chance=None,
             chance_note="—؛ IELTS ۷٫۵ و احراز حداقل‌ها لازم",
             entry="حد B+ / ۷۸٪ کل یا معادل؛ زمینهٔ کمی و درس‌های پیشرفته؛ IELTS ۷٫۵، Writing/Speaking حداقل ۷؛ سه referee، حداقل دو علمی.",
             why="بودجه دیگر عامل حذف نیست، اما با سناریوی زبان ۷٫۰ شرط زبان احراز نمی‌شود. آمار ۱۲، حسابان ۱۰ و الگوریتم ۱۱٫۵ در ارزیابی باقی‌اند؛ از نظر تناسب بسیار بلندپروازانه/مشروط است. co-op شغل تضمینی نیست؛ درصد جدید ساخته نشده.",
             sources=["waterloo_ds","waterloo_ds_lang","waterloo_fees"], fee_cad=Decimal("6272")*9,
             budget_note="۹ درس نرخ 2026/27؛ incidental و co-op جدا. هزینه صرفاً اطلاع، نه حذف مالی.",
             priority="مشروط به زبان و بررسی جدی پایهٔ کمی", duration="۱۶ ماه", scenario_added=True),
        dict(group="کامپیوتر کاربردی", uni="Toronto", program="MScAC — Computer Science / AI / Data Science",
             route="coursework + applied-research internship؛ حرفه‌ای، نه thesis معمولی", chance=None,
             chance_note="—؛ بسیار رقابتی، درصد معتبر نداریم",
             entry="حد B+ معادل در سال آخر؛ پیش‌نیازِ گرایش، زبان SGS و سه referee. درخواست Fall 2027 از ۱ Oct تا ۱ Dec 2026 باز است.",
             why="معدل سال آخر ۱۷٫۰۳ یک نکتهٔ مثبت است، نه اثبات معادل B+. برای تو CS نسبت به AI/DS کمی تناسب بهتری دارد: OS ۱۹ و DB ۱۶٫۵ مثبت، الگوریتم ۱۱٫۵ ریسک است. AI/DS نیز زمینهٔ جدی کمی می‌خواهند. گرایش‌ها یک خانوادهٔ برنامه‌اند؛ چند فرصت مستقل پذیرش حساب نشده‌اند. هیچ نرخ فردی تازه‌ای ادعا نمی‌شود.",
             sources=["toronto_mscac","toronto_mscac_sgs"], fee_cad=Decimal("90050"),
             budget_note="برآورد tuition/fees ورودی ۲۰۲۶، نه رقم نهایی ۲۰۲۷؛ فقط برای اطلاع. با فرض بودجهٔ حل‌شده، عامل حذف نیست.",
             priority="اپلای بلندپروازانه؛ ترجیح CS در میان گرایش‌ها", duration="۱۶ ماه؛ شامل internship پژوهشی کاربردی", scenario_added=True),
    ]
    prior_keys={
        ("Waterloo","MEng ECE — Software"),
        ("Western","MEng ECE — Software"),
        ("Waterloo","MDSAI"),
        ("Toronto","MScAC — AI / Data Science"),
    }
    assert len(ITEMS)==20 and len(EXCLUDED)==6
    assert prior_keys.issubset({(x['uni'],x['program']) for x in EXCLUDED})
    ITEMS.extend(additions)
    EXCLUDED[:]=[x for x in EXCLUDED if (x['uni'],x['program']) not in prior_keys]
    for x in ITEMS:
        x['budget_note']=x['budget_note'].replace("داخل بودجه فرض نشود.","صرفاً برای برنامه‌ریزی هزینه تأیید شود؛ محدودیت مالی عامل حذف نیست.")
        x['budget_note']=x['budget_note'].replace("گام۱ پایه زیر سقف است، ولی ","")
        x['priority']=x['priority'].replace("؛ هزینه نامعلوم","")
    updates={
        "دامنه":"فقط کانادا بدون کبک، ورودی سپتامبر ۲۰۲۷؛ امنیت حذف؛ در این نسخه بودجه حل‌شده فرض شده و هیچ ردیفی به دلیل سقف مالی حذف نمی‌شود.",
        "بودجه":"سناریوی فرضیِ بودجهٔ کافی است، نه تغییر سقف واقعی بودجه در فایل مهاجرت قبلی. شهریه/تمکن فقط اطلاعات‌اند. مانع علمی، زبان، مدرک، عرضهٔ برنامه و نیاز به استاد همچنان اعمال می‌شوند.",
        "موارد بدون درصد":"UVic/Manitoba استادمحورند؛ برای MScAC آمار فردی معتبر نداریم؛ MDSAI در سناریوی IELTS ۷ شرط ۷٫۵ را ندارد. به‌جای حدس جدید، دلیل خالی بودن هر ردیف نوشته شده است. بازه‌های قدیمی با افزایش بودجه بالا نرفته‌اند.",
    }
    ASSUMPTIONS[:]=[(a,updates.get(a,b)) for a,b in ASSUMPTIONS]
    ASSUMPTIONS.append(("گرایش‌های تورنتو","MScAC یک برنامه با چند concentration است؛ CS/AI/DS را سه اپلای مستقل یا سه احتمال مستقل حساب نکن. B+ سال آخر و پیش‌نیاز همان concentration باید با معیار رسمی تطبیق داده شود."))
    ASSUMPTIONS.append(("نسخهٔ حفظ‌شده","نسخهٔ دارای سقف حدود $45,000 و فایل‌های اصلی مهاجرت/کارنامه حفظ شده‌اند؛ این خروجی فایل مستقلِ No-Budget است."))
    assert len(ITEMS)==24 and len(EXCLUDED)==2


def chance_for_row(x):
    return x.get('chance_note') or chance_text(x['chance'])


def no_budget_markdown():
    L=["# فهرست شخصی کانادا — سناریوی بودجهٔ حل‌شده، بدون امنیت", "",
       "**۱ اکتبر ۲۰۲۶ | ورودی سپتامبر ۲۰۲۷ | ۲۴ برنامه/خانوادهٔ برنامه**", "",
       "> **تغییر این نسخه:** سقف مالی از فیلتر حذف شده است. MEng نرم‌افزار واترلو، MEng نرم‌افزار وسترن، MDSAI واترلو و MScAC تورنتو برگشته‌اند. پول داشتن، حداقل معدل، زبان، پیش‌نیاز یا رقابت را برطرف نمی‌کند.", "",
       "> **درصدها:** بازه‌های ستاره‌دار قبلی فقط قضاوت خام و کالیبره‌نشده‌اند؛ احتمال واقعی محاسبه نشده و با تغییر بودجه افزایش داده نشده‌اند. برای MScAC و MDSAI عدد تازهٔ بی‌پشتوانه ساخته نشده است.", "",
       "> **سناریوی زبان:** همچنان IELTS Academic فرضی ۷٫۰ با هر band ≥۶٫۵؛ نمره واقعی ارائه نشده. MDSAI شرط جداگانهٔ ۷٫۵ overall و W/S ≥۷ دارد؛ در فهرست مشروط آمده، نه به‌عنوان احراز‌شده.", "",
       "## روش و پروفایل", "",
       "رتبهٔ موضوعی: CS در THE ۲۰۲۶؛ رتبهٔ کلی: QS ۲۰۲۷. هر حوزه اول رتبهٔ CS، سپس دانشگاه؛ رتبهٔ خود مدرک یا رتبهٔ مستقل Software/Digital Media ادعا نمی‌شود. مقیاس کارنامه همچنان ۲۰: کل ۱۵٫۷۷، دو سال آخر ۱۶٫۹۲، سال آخر ۱۷٫۰۳، سه سال آخر ۱۶٫۱۹. همهٔ اعداد مالی USD و صرفاً برای اطلاع‌اند.", "",
       "## چهار مورد اضافه‌شده", "",
       "| دانشگاه | برنامه | رتبهٔ CS / THE 2026 | رتبهٔ کلی / QS 2027 | ارزیابی / شرط اصلی |", "|---|---|---:|---:|---|"]
    for x in sorted([x for x in ITEMS if x.get('scenario_added')],key=sort_key):
        a,b=UNI[x['uni']]
        L.append(f"| {x['uni']} | {x['program']} | {fmt_rank(a)} | {fmt_rank(b)} | {chance_for_row(x)}؛ {x['priority']} |")
    for group in GROUPS:
        L += ["",f"## {group}","", "| دانشگاه | برنامه / نوع | CS / THE 2026 | دانشگاه / QS 2027 | بازهٔ قبلی یا وضعیت | هزینهٔ پایه USD، فقط اطلاع |", "|---|---|---:|---:|---:|---:|"]
        rs=sorted([x for x in ITEMS if x['group']==group],key=sort_key)
        for x in rs:
            a,b=UNI[x['uni']]
            L.append(f"| {x['uni']} | {x['program']} — {x['route']} | {fmt_rank(a)} | {fmt_rank(b)} | {chance_for_row(x)} | {budget_text(x)} |")
        for x in rs:
            links=" ".join(f"[{i+1}]({SOURCES[k][1]})" for i,k in enumerate(x['sources']))
            L += ["",f"### {x['uni']} — {x['program']}","",f"- **نوع/مدت:** {x['route']}؛ {x['duration']}.",f"- **شرط ورود:** {x['entry']}",f"- **تناسب/ریسک:** {x['why']}",f"- **وضعیت:** {chance_for_row(x)}؛ {x['priority']}.",f"- **هزینه، فقط اطلاع:** {budget_text(x)}؛ {x['budget_note']}",f"- **منابع رسمی:** {links}"]
    L += ["", "## هنوز نیازمند رفع مانع، نه مانع بودجه", "",
          "| دانشگاه / برنامه | دلیل باقی‌ماندن خارج از سبد اصلی |", "|---|---|"]
    for x in EXCLUDED:
        L.append(f"| {x['uni']} / {x['program']} | {x['reason']} |")
    L += ["", "- **UBC / SFU:** کف معمول معدل کل ایران ۱۶/۲۰ در برابر ۱۵٫۷۷ تو، بدون تأیید استثنا هنوز مانع است؛ بودجه این شرط را تغییر نمی‌دهد.",
          "- **پایان‌نامه‌ای خالص:** با ترجیح درسی تو خودکار به سبد اصلی اضافه نشده‌اند.",
          "- **امنیت / کبک:** همچنان از توصیه‌ها حذف‌اند؛ نمرهٔ درس امنیت از معدل واقعی پاک نشده است.", "",
          "## ترتیب بررسی علمی من با حذف سقف مالی", "",
          "مسیرهای نرم‌افزاریِ Western MEng، McMaster MEng، Queen's/Calgary MEng و Memorial Software همچنان با قوت درس‌های نرم‌افزارت متناسب‌اند. Waterloo MEng Software گزینهٔ بلندپروازانهٔ مرتبط است. MScAC تورنتو با ترجیح concentrationِ CS، و MDSAI واترلو، انتخاب‌های بسیار رقابتی/مشروط‌اند؛ بودجهٔ کافی آن‌ها را گزینهٔ مطمئن نمی‌کند.", "",
          "برای MScAC تورنتو، معیار رسمی B+ در سال آخر است؛ ۱۷٫۰۳ سال آخر یک قوت قابل اشاره است، نه معادل‌سازی قطعی. در CS، OS/Databases/Algorithms مهم‌اند: ۱۹ و ۱۶٫۵ مثبت، ۱۱٫۵ ریسک. در AI/DS، زمینهٔ کمیِ بیشتر لازم است. MScAC یک خانوادهٔ برنامه است؛ گرایش‌هایش فرصت‌های مستقلِ آماری تلقی نشده‌اند.", "",
          "**چرخهٔ رسمی Fall 2027 تورنتو:** درخواست‌ها از ۱ اکتبر ۲۰۲۶ باز شده و تا ۱ دسامبر ۲۰۲۶، ساعت ۱۱:۵۹ شب ET، بسته می‌شوند. این تاریخ از صفحهٔ رسمی MScAC است، نه پیش‌بینی عمومی.", "",
          "## فرض‌ها و محدودیت‌ها", ""]
    for a,b in ASSUMPTIONS:
        L.append(f"- **{a}:** {b}")
    L += ["", "## منابع", "", "صفحات رسمی مرتبط با چهار مورد برگشتی دوباره در ۱ اکتبر ۲۰۲۶ بررسی شدند؛ وجود اطلاعات مالی قدیمی، ادعای نرخ دقیق ۲۰۲۷ نیست.", ""]
    for k,(title,url) in SOURCES.items():
        L.append(f"- **{title}** — {url}")
    (ROOT/f'{OUTPUT_STEM}.md').write_text("\n".join(L)+"\n",encoding='utf-8')


def fmt_rank(x): return x or "— / در مرجع تأیید نشد"
FA_TRANS = str.maketrans("0123456789.", "۰۱۲۳۴۵۶۷۸۹٫")
def fa(x): return str(x).translate(FA_TRANS)
def rank_key(x):
    if not x: return (inf, inf)
    nums=re.findall(r"\d+",x)
    return (int(nums[0]),int(nums[-1]))
def sort_key(x):
    a,b=UNI[x['uni']]
    return (rank_key(a),rank_key(b),x['uni'],x['program'])
def chance_text(x):
    if x is None: return "—؛ استادمحور/احراز شرط"
    return f"{fa(x[0])}–{fa(x[1])}٪*"
def cash(v): return f"${int(v):,}"
def baseline(fee):
    if fee is None: return None
    return ((fee+PROOF_CAD)/FX).quantize(Decimal('1'),rounding=ROUND_HALF_UP)
def budget_text(x):
    b=baseline(x['fee_cad'])
    if b is None: return "—؛ هزینه استخراج‌نشده (اطلاع)" if IGNORE_BUDGET else "—؛ تأیید هزینه لازم"
    return f"≈ {cash(b)} پایهٔ 2026/27"

def verify_grades():
    with (ROOT/'Grades-extracted.csv').open(encoding='utf-8-sig',newline='') as f:
        rr=list(csv.DictReader(f))
    ps=[r for r in rr if r['وضعیت']=='گذرانده' and int(r['تعداد واحد'])>0]
    def avg(xs):
        cr=sum(int(r['تعداد واحد']) for r in xs)
        p=sum(Decimal(r['نمره از 20'])*int(r['تعداد واحد']) for r in xs)
        return cr,(p/cr).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
    assert avg(ps)==(140,Decimal('15.77'))
    assert avg([r for r in ps if r['سال و نیمسال'].startswith(('1403-1404','1404-1405'))])==(77,Decimal('16.92'))
    assert avg([r for r in ps if r['سال و نیمسال'].startswith('1404-1405')])==(38,Decimal('17.03'))
    assert avg([r for r in ps if r['سال و نیمسال'].startswith(('1402-1403','1403-1404','1404-1405'))])==(111,Decimal('16.19'))
    return ps

def markdown():
    if IGNORE_BUDGET:
        return no_budget_markdown()
    L=["# فهرست شخصی ارشدهای درسی/پروژه‌ای کانادا — بدون امنیت", "", "**۱ اکتبر ۲۰۲۶ | ورودی هدف سپتامبر ۲۰۲۷ | کانادا بدون کبک**", "",
       "> **هشدار دربارهٔ درصدها:** اعداد ستاره‌دار فقط بازهٔ قضاوت شخصیِ بسیار خام و کالیبره‌نشده‌اند. احتمال واقعیِ پذیرش از دادهٔ تاریخیِ هم‌پروفایل محاسبه نشده است؛ هیچ مدل آماری یا تضمینی وجود ندارد. حتی با کارنامهٔ کامل، IELTS، SOP، reference، معادل‌سازی، ظرفیت و تصمیم ۲۰۲۷ را نداریم. این بازه‌ها فقط برای مقایسهٔ سبدند، نه پیش‌بینی مالی؛ واقعیت ممکن است بیرون آن‌ها باشد.", "",
       "> **فرض زبان:** IELTS Academic برابر ۷٫۰ با هر band حداقل ۶٫۵؛ این نمره هنوز کسب/ارائه نشده است. تمام برآوردها مشروط به احراز حداقل رسمیِ همان برنامه‌اند. برای موارد استادمحور/مانع سخت، به‌جای عددسازی خانه خالی است.", "",
       "## روش رتبه‌بندی", "", "رتبهٔ حوزه: **THE Computer Science 2026**؛ رتبهٔ کلی: **QS 2027**. این رتبه‌ها متعلق به دانشگاه/حوزه‌اند، نه خود مدرک MEng/MSc. برای Software، AI و Multimedia شاخص CS مشترک است؛ برای Digital Media فقط یک proxy مرتبط است و ادعای رتبهٔ مستقل آن رشته نیست. در هر حوزه اول CS و سپس رتبهٔ کلی؛ band مشترک رتبهٔ دقیقِ ساختگی نمی‌گیرد. [1](https://www.timeshighereducation.com/student/best-universities/best-universities-canada-computer-science-degrees)", "",
       "## مبنای پروفایل", "", "کارنامه: کل **۱۵٫۷۷**؛ دو سال آخر **۱۶٫۹۲**؛ سال آخر **۱۷٫۰۳**؛ سه سال آخر **۱۶٫۱۹**. قوت اصلی نرم‌افزار **۱۹٫۲۵**، پروژه **۱۹٫۵**، OS **۱۹**، برنامه‌نویسی و جبر **۱۸** است؛ حسابان/آمار/الگوریتم ریسک‌اند. امنیت مطابق خواستهٔ تو از توصیه‌ها کنار گذاشته شده است، نه از معدل رسمی.", "",
       "پروژهٔ املاک و ETL سینما طبق READMEهای public، برای software/product/applied AI قابل استفاده‌اند. API integration و توصیه‌گرِ exact-genre را آموزش مدل یا مقالهٔ ML حساب نکرده‌ام. سه‌سال سابقهٔ IT را سه‌سال توسعهٔ نرم‌افزار تلقی نمی‌کنم؛ README شغل فعلی را از Mar 2025 ذکر می‌کند و سابقهٔ قبلی برای ۳ سال نیازمند مستند است.", ""]
    for group in GROUPS:
        L += [f"## {group}", "", "| دانشگاه | برنامه / نوع | رتبه CS / THE 2026 | رتبه کلی / QS 2027 | بازهٔ شخصیِ مشروط* | گام۱ پایه، USD |", "|---|---|---:|---:|---:|---:|"]
        rs=sorted([x for x in ITEMS if x['group']==group],key=sort_key)
        for x in rs:
            a,b=UNI[x['uni']]
            L.append(f"| {x['uni']} | {x['program']} — {x['route']} | {fmt_rank(a)} | {fmt_rank(b)} | {chance_for_row(x)} | {budget_text(x)} |")
        L.append("")
        for x in rs:
            links=" ".join(f"[{i+1}]({SOURCES[k][1]})" for i,k in enumerate(x['sources']))
            # Links are report-local source labels, not claimed official probabilities.
            L += [f"### {x['uni']} — {x['program']}", "", f"- **نوع/مدت:** {x['route']}؛ {x['duration']}.", f"- **شرط ورود:** {x['entry']}", f"- **قضاوت تناسب:** {x['why']}", f"- **بودجه:** {budget_text(x)}؛ {x['budget_note']}", f"- **اولویت:** {x['priority']}؛ بازهٔ خام {chance_for_row(x)}.", f"- **منابع رسمی:** {links}", ""]
    L += ["## خارج از سبد اصلی فعلی", "", "بودجه با شانس پذیرش یکی نیست؛ خروج مالی را احتمال پذیرش صفر معنا نکن. درصدهای موجود در این بخش فقط نگاه علمیِ مشروطند و برنامه را قابل تأمین نمی‌کنند.", "", "| دانشگاه / برنامه | رتبه CS | رتبه کلی | بازهٔ خام علمی* | گام۱ پایه USD | علت خروج |", "|---|---:|---:|---:|---:|---|"]
    for x in sorted(EXCLUDED,key=sort_key):
        a,b=UNI[x['uni']]
        L.append(f"| {x['uni']} / {x['program']} | {fmt_rank(a)} | {fmt_rank(b)} | {chance_for_row(x)} | {budget_text(x)} | {x['reason']} |")
    L += ["", "UBC و SFU نیز با معدل کل ۱۵٫۷۷ در برابر کف معمولِ رسمی ایرانِ ۱۶/۲۰، بدون تأیید استثنا در سبد اصلی نیستند؛ هیچ برنامهٔ کبک یا امنیت پیشنهاد نشده است.", "", "## بودجه: عدد پایه با هزینهٔ کامل فرق دارد", "",
          f"تمکنِ فعلی یک نفر، با فرض تبدیل همین تحلیل، حدود **{cash((PROOF_CAD/FX).quantize(Decimal('1'),rounding=ROUND_HALF_UP))}** است. معیار جدول، کل شهریهٔ دوره + یک سال تمکن است، نه مجموع زندگی تمام دوره. هزینهٔ اجباری، سفر، ویزا، ارز و ۲۰۲۷ هنوز مهم‌اند؛ درآمد دانشجویی/کارآموزی تضمین نیست.", "",
          "**اصلاح مالی نسبت به پیشنهاد علمی قبلی:** MEng واترلو در ۴ term عادی حدود **$65,573** و MEng وسترن در ۳ term حدود **$49,753** شهریه+تمکن پایه دارند، قبل ancillary. پس با حدود $45,000 گزینهٔ اصلی نیستند. Multimedia آلبرتا با نرخ رسمی 2026/27 و تمام ۳۶ واحد حدود **$39,017** پایه است؛ نرخ قدیمی ۲۰۲۲ در محاسبه استفاده نشده است.", "",
          "## ترتیب عملیِ پیشنهادی من، متفاوت از ترتیب رنک", "", "۱. Memorial MASc Software Engineering؛ ۲. Alberta Multimedia با portfolio مرتبط؛ ۳. McMaster MEng Computing and Software، فقط پس از تأیید مالی؛ ۴. UNB MCSC و Memorial MSc CS all-course؛ ۵. Regina MEng SSE / Queen's MEng / Calgary Software، پس از بررسی هزینه و ورود. مسیرهای AI/Data برای تو گزینهٔ دوم‌اند، نه جایگزین خودکارِ نرم‌افزار. اعدادِ شانس نباید این ترتیب را به یک پیش‌بینی آماری تبدیل کنند.", "", "## فرض‌ها و محدودیت‌ها", ""]
    for title,text in ASSUMPTIONS: L += [f"- **{title}:** {text}"]
    L += ["", "## منابع و تاریخ", "", "منابع در ۱ اکتبر ۲۰۲۶ بررسی شدند. وجود برنامه و قواعد فعلی به‌معنی تضمین همهٔ شرایط ورودی ۲۰۲۷ نیست. هیچ مقدارِ هزینهٔ استخراج‌نشده، صفر یا عدد aggregator معرفی نشده است.", ""]
    for k,(title,url) in SOURCES.items(): L.append(f"- **{title}** — {url}")
    (ROOT/'Canada-Course-Shortlist.md').write_text("\n".join(L)+"\n",encoding='utf-8')

NAVY="18394A"; TEAL="236D77"; LIGHT="F2F7F8"; AMBER="FFF0CB"; RED="FCE2E2"; WHITE="FFFFFF"
def base_style(ws, widths):
    ws.sheet_view.rightToLeft=True
    ws.sheet_view.showGridLines=False
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.outlinePr.summaryRight=False

def banner(ws,title,n=8):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
    ws.cell(1,1,title).font=Font(name=FONT,size=20,bold=True,color=WHITE)
    ws.cell(1,1).fill=PatternFill('solid',fgColor=NAVY)
    ws.cell(1,1).alignment=Alignment(horizontal='right',vertical='center',readingOrder=2)
    ws.row_dimensions[1].height=39
    texts=[("کانادا بدون کبک | ورودی ۲۰۲۷ | امنیت حذف | بودجه فیلتر نیست" if IGNORE_BUDGET else "کانادا بدون کبک | ورودی ۲۰۲۷ | امنیت حذف | درسی/پروژه‌ای"),
           "هشدار: درصدهای * فقط نظر بسیار خام، مشروط و کالیبره‌نشده‌اند؛ نه آمار، مدل معتبر یا تضمین. IELTS فرضی ۷٫۰، هر band ≥۶٫۵.",
           "ترتیب: رتبهٔ CS از THE 2026، سپس رتبهٔ کلی QS 2027؛ رتبهٔ خود برنامه نیست. هزینه‌ها USD و پایهٔ 2026/27؛ ۲۰۲۷ و ancillary جدا."]
    if IGNORE_BUDGET:
        texts[1]="درصدهای * نظر خامِ قبلی‌اند و افزایش نیافته‌اند؛ IELTS فرضی ۷٫۰، هر band ≥۶٫۵. MDSAI: شرط جداگانهٔ ۷٫۵ و W/S ≥۷."
        texts[2]="رتبهٔ CS: THE 2026، کلی: QS 2027؛ رتبهٔ خود برنامه نیست. همهٔ هزینه‌های USD صرفاً اطلاع‌اند؛ هیچ حذف مالی نداریم."
    for r,text in enumerate(texts,2):
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=n)
        c=ws.cell(r,1,text); c.font=Font(name=FONT,size=12,bold=r==3,color=NAVY)
        c.fill=PatternFill('solid',fgColor=AMBER if r==3 else LIGHT)
        c.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True,readingOrder=2)
        ws.row_dimensions[r].height=33 if r!=3 else 42

def table_sheet(wb,title,records,excluded=False):
    ws=wb.create_sheet(title)
    base_style(ws,[5,20,46,15,15,22,24,62])
    banner(ws,f"{title} — رنک و برآوردِ شخصی مشروط")
    headers=['#','دانشگاه','برنامه / نوع','CS / THE 2026','دانشگاه / QS 2027','نظر شخصی، ٪*','شهریه + تمکن، USD','چرا / شرط اصلی / وضعیت هزینه']
    for j,t in enumerate(headers,1):
        c=ws.cell(6,j,t); c.font=Font(name=FONT,size=12,bold=True,color=WHITE)
        c.fill=PatternFill('solid',fgColor=TEAL)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True,readingOrder=2)
    ws.row_dimensions[6].height=40
    for i,x in enumerate(sorted(records,key=sort_key),1):
        r=6+i; a,b=UNI[x['uni']]
        detail=x['reason'] if excluded else f"{x['entry']}\n{x['why']}\n{'هزینه، فقط اطلاع' if IGNORE_BUDGET else 'بودجه'}: {x['budget_note']}"
        prog=x['program'] if excluded else f"{x['program']}\n{x['route']} | {x['duration']}"
        vals=[i,x['uni'],prog,fmt_rank(a),fmt_rank(b),chance_for_row(x),budget_text(x),detail]
        for j,v in enumerate(vals,1):
            c=ws.cell(r,j,v); c.font=Font(name=FONT,size=12,color=NAVY,bold=j in [2,6])
            c.fill=PatternFill('solid',fgColor=WHITE if i%2 else LIGHT)
            c.alignment=Alignment(horizontal='right' if j in [3,7,8] else 'center',vertical='center',wrap_text=True,readingOrder=2)
            c.border=Border(bottom=Side(style='hair',color='D9E4E9'))
        ws.cell(r,6).fill=PatternFill('solid',fgColor=AMBER)
        ws.cell(r,6).comment=Comment("قضاوت شخصی کالیبره‌نشده است؛ احتمال واقعی یا confidence interval محاسبه نشده. مشروط به همهٔ حداقل‌ها، IELTS فرضی و مدارک خوب؛ نتیجهٔ واقعی ممکن است بیرون بازه باشد. برای تصمیم مالی یا احتمال مجموع اپلای استفاده نشود.","یادداشت روش")
        if x['fee_cad'] is not None:
            value=baseline(x['fee_cad'])
            ws.cell(r,7).fill=PatternFill('solid',fgColor=LIGHT if IGNORE_BUDGET else (RED if value>BUDGET_USD else 'E5F0EC'))
            ws.cell(r,7).comment=Comment("فقط کل شهریهٔ پایهٔ 2026/27 + یک سال تمکن، با فرض USD/CAD = 1.419؛ ancillary، سفر/ویزا، ارز و تغییرات ۲۰۲۷ جداست. زیر سقف بودن این عدد، تأیید کامل مالی نیست.","یادداشت بودجه")
        else:
            ws.cell(r,7).fill=PatternFill('solid',fgColor=LIGHT if IGNORE_BUDGET else AMBER)
        ws.cell(r,3).hyperlink=SOURCES[x['sources'][0]][1]
        ws.cell(r,3).comment=Comment("لینک صفحهٔ رسمی برنامه؛ همهٔ منابع در شیت منابع آمده‌اند.","منبع")
        ws.row_dimensions[r].height=150 if excluded else 180
    if records:
        tab=Table(displayName=f"List{len(wb.worksheets)}",ref=f"A6:H{6+len(records)}")
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=False,showColumnStripes=False)
        ws.add_table(tab)
    ws.freeze_panes='D7'; ws.print_title_rows='1:6'; ws.print_area=f'A1:H{max(6,6+len(records))}'
    return ws

def workbook(grades):
    wb=Workbook(); wb.remove(wb.active)
    for g in GROUPS:
        table_sheet(wb,g,[x for x in ITEMS if x['group']==g])
    review_sheet='نیازمند رفع مانع' if IGNORE_BUDGET else 'خارج سبد اصلی'
    table_sheet(wb,review_sheet,EXCLUDED,True)
    ws=wb.create_sheet('فرض‌ها و روش'); base_style(ws,[29,112]); banner(ws,'فرض‌ها و حدود اعتبار برآوردها',2)
    ws.append([])
    for r,(a,b) in enumerate(ASSUMPTIONS,6):
        ws.cell(r,1,a); ws.cell(r,2,b)
        for c in ws[r][:2]:
            c.font=Font(name=FONT,size=13,color=NAVY,bold=c.column==1)
            c.alignment=Alignment(horizontal='right',vertical='center',readingOrder=2,wrap_text=True)
            c.fill=PatternFill('solid',fgColor=LIGHT if r%2==0 else WHITE)
        ws.row_dimensions[r].height=82 if a not in ['قانون درصدها','رتبهٔ موضوعی','سابقه کار'] else 108
    ws.freeze_panes='B6'
    ws=wb.create_sheet('نمرات مرتبط'); base_style(ws,[42,13,17,26]); banner(ws,'شواهد نمره — بدون تغییر کارنامه',4)
    names={'مهندسی نرم‌افزار','پروژه کارشناسی','طراحی زبان‌های برنامه‌سازی','سیستم‌های عامل','برنامه‌سازی وب','مبانی برنامه‌سازی','برنامه‌سازی پیشرفته','ساختمان داده و الگوریتم‌ها','اصول طراحی کامپایلر','طراحی پایگاه داده‌ها','جبر خطی','ساختمان‌های گسسته','ریاضی عمومی 1','ریاضی عمومی 2','آمار و احتمالات مهندسی','معادلات دیفرانسیل','طراحی الگوریتم‌ها','هوش مصنوعی','مبانی یادگیری ماشین','گرافیک کامپیوتری','معماری کامپیوتر','طراحی سیستم‌های دیجیتال','طراحی VLSI'}
    for j,v in enumerate(['درس','واحد','نمره از ۲۰','زمان'],1): ws.cell(6,j,v)
    r=7
    for x in grades:
        if x['نام درس'] not in names: continue
        vals=[x['نام درس'],int(x['تعداد واحد']),float(x['نمره از 20']),x['سال و نیمسال']]
        for j,v in enumerate(vals,1):
            c=ws.cell(r,j,v); c.font=Font(name=FONT,size=13,color=NAVY)
            c.alignment=Alignment(horizontal='right' if j==1 else 'center',vertical='center',readingOrder=2)
        ws.cell(r,3).number_format='0.00'; ws.row_dimensions[r].height=29; r+=1
    ws.freeze_panes='B7'; ws.auto_filter.ref=f'A6:D{r-1}'
    ws=wb.create_sheet('منابع'); base_style(ws,[22,80,110]); banner(ws,'منابع رسمی و مدارک بررسی‌شده',3)
    for j,v in enumerate(['کلید','منبع / کاربرد','URL'],1): ws.cell(6,j,v)
    for r,(k,(title,url)) in enumerate(SOURCES.items(),7):
        for j,v in enumerate([k,title,url],1):
            c=ws.cell(r,j,v); c.font=Font(name=FONT,size=12,color=NAVY)
            c.alignment=Alignment(horizontal='right' if j==2 else 'left',vertical='center',wrap_text=True,readingOrder=2 if j==2 else 1)
        ws.cell(r,3).hyperlink=url; ws.row_dimensions[r].height=45
    ws.freeze_panes='C7'; ws.auto_filter.ref=f'A6:C{6+len(SOURCES)}'
    for sheet_name in ['نمرات مرتبط', 'منابع']:
        sh = wb[sheet_name]
        for c in sh[6]:
            if c.value is None:
                continue
            c.font = Font(name=FONT, size=12, bold=True, color=WHITE)
            c.fill = PatternFill('solid', fgColor=TEAL)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
        sh.row_dimensions[6].height = 36
    wb.active=0
    dest=ROOT/f'{OUTPUT_STEM}.xlsx'; wb.save(dest)
    test=load_workbook(dest)
    assert test.sheetnames[:4]==GROUPS
    assert len(ITEMS)==(24 if IGNORE_BUDGET else 20)
    for g in GROUPS:
        assert test[g].max_row==6+len([x for x in ITEMS if x['group']==g])
    assert test[review_sheet].max_row==6+len(EXCLUDED)
    assert all(test[s].sheet_view.rightToLeft for s in test.sheetnames)
    return dest


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ignore-budget', action='store_true', help='Generate a separate 24-row No-Budget scenario, keeping academic/language filters.')
    args=parser.parse_args()
    configure_scenario(args.ignore_budget)
    grades=verify_grades()
    assert all(not re.search(r'security|cyber|امنیت',x['program'],re.I) for x in ITEMS)
    assert all(x['uni'] in UNI for x in ITEMS+EXCLUDED)
    for x in ITEMS+EXCLUDED:
        assert all(k in SOURCES for k in x['sources'])
        if x['chance'] is not None:
            lo,hi=x['chance']; assert 0<=lo<hi<=100
    markdown(); dest=workbook(grades)
    print(f'Created {dest.name} and {OUTPUT_STEM}.md: {len(ITEMS)} candidates, {len(EXCLUDED)} excluded/comparison rows.')
    print('All admission ranges are manual, subjective and uncalibrated; not empirical probabilities.')

if __name__=='__main__': main()
